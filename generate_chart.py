import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Alignment, Font, Border, Side
from openpyxl.utils import get_column_letter

# Define project data
phases = [
    "1. Project Setup & Scoping",
    "2. Discovery & Analysis",
    "3. Framework Development",
    "4. Testing & Validation",
    "5. Evaluation & Optimization",
    "6. Documentation & Report Finalization"
]
start_weeks = [1, 2, 4, 7, 9, 10]
end_weeks = [1, 3, 6, 8, 9, 11]

# Workbook setup
wb = Workbook()
ws = wb.active
ws.title = "Gantt Chart"

# Header row
headers = ["Phase", "Start Week", "End Week"] + [f"Week {i}" for i in range(1, 12)]
ws.append(headers)

# Styles
header_font = Font(bold=True, color="FFFFFF")
header_fill = PatternFill(start_color="2E4053", end_color="2E4053", fill_type="solid")
fill_colors = ["4A90E2", "50C878", "E67E22", "AF7AC5", "E74C3C", "17A589"]
border = Border(left=Side(style="thin"), right=Side(style="thin"),
                top=Side(style="thin"), bottom=Side(style="thin"))

# Apply header styles
for col_num, cell in enumerate(ws[1], 1):
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = Alignment(horizontal="center", vertical="center")
    cell.border = border
    ws.column_dimensions[get_column_letter(col_num)].width = 13

# Populate data
for i, phase in enumerate(phases):
    start = start_weeks[i]
    end = end_weeks[i]
    color = PatternFill(start_color=fill_colors[i % len(fill_colors)],
                        end_color=fill_colors[i % len(fill_colors)], fill_type="solid")

    row = [phase, start, end] + [""] * 11
    ws.append(row)

    # Fill timeline cells with color
    for col in range(start + 3, end + 4):  # +3 offset for first 3 columns
        ws.cell(row=i + 2, column=col).fill = color

# Alignment & borders for all cells
for row in ws.iter_rows(min_row=2, max_row=len(phases) + 1, min_col=1, max_col=14):
    for cell in row:
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border

# Add title
ws.merge_cells("A8:N8")
title_cell = ws["A8"]
title_cell.value = "Gantt Chart: Semi-Automated Migration from Windows 11 to Linux Mint"
title_cell.font = Font(size=14, bold=True, color="FFFFFF")
title_cell.fill = PatternFill(start_color="1F618D", end_color="1F618D", fill_type="solid")
title_cell.alignment = Alignment(horizontal="center", vertical="center")

# Save file
file_path = ".\\docs\\Decorated_Windows_to_Linux_Gantt.xlsx"
wb.save(file_path)

file_path
